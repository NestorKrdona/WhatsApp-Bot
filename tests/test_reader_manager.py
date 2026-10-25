import pytest
from unittest import mock
import pandas as pd
from src.modules.reader_manager import ReaderHandler
from openpyxl import Workbook

@pytest.fixture
def workbook_and_worksheet():
    """Fixture para crear un libro y hoja de trabajo de Excel"""
    wb = Workbook()
    ws = wb.active
    return wb, ws

@pytest.fixture
def mock_cell():
    return mock.MagicMock()

@pytest.fixture
def mock_excel_read():
    """Simula pd.read_excel para evitar tener que leer un archivo real."""
    with mock.patch("pandas.read_excel") as mock_read_excel:
        yield mock_read_excel

def test_excel_sum_column_no_filter(mock_excel_read):
    """Verifica que funcione la suma de columnas sin filtro"""
    data = {
        "A": ["ACTIVO", "INACTIVO", "ACTIVO", "ACTIVO", "INACTIVO"],
        "B": [10, 20, 30, 40, 50]
    }
    df = pd.DataFrame(data)
    mock_excel_read.return_value = df
    obj = ReaderHandler()
    result = obj.excel_sum_column("fake_path.xlsx", 0, 1)
    assert result == 150

def test_excel_sum_column_with_filter(mock_excel_read):
    """Verifica que funcione la suma de columnas con filtro de ACTIVO"""
    data = {
        "A": ["ACTIVO", "INACTIVO", "ACTIVO", "ACTIVO", "INACTIVO"],
        "B": [10, 20, 30, 40, 50]
    }
    df = pd.DataFrame(data)
    mock_excel_read.return_value = df  
    obj = ReaderHandler()
    result = obj.excel_sum_column("fake_path.xlsx", 0, 1, filter_arg="ACTIVO")
    assert result == 80

def test_excel_sum_column_without_numbers(mock_excel_read):
    """Verifica que funcione solo para las columnas con datos sumables"""
    data = {
        "A": ["ACTIVO", "ACTIVO", "ACTIVO", "ACTIVO", "ACTIVO"],
        "B": [10, "texto", 30, None, 50]
    }
    df = pd.DataFrame(data)
    mock_excel_read.return_value = df 
    obj = ReaderHandler()
    result = obj.excel_sum_column("fake_path.xlsx", 0, 1)
    assert result == 90


def test_excel_sum_column_empty_column(mock_excel_read):
    """Verifica el caso de prueba en caso de que la columna no tenga datos"""
    data = {
        "A": ["ACTIVO", "ACTIVO", "ACTIVO"],
        "B": [None, None, None]
    }
    df = pd.DataFrame(data)
    mock_excel_read.return_value = df
    obj = ReaderHandler()
    result = obj.excel_sum_column("fake_path.xlsx", 0, 1)
    assert result == 0

def test_set_cell_merged(mock_cell):
    merged_cell_mock = mock_cell
    merged_cell_mock.coordinate = "A1"
    merged_range_mock = mock.MagicMock(min_row=1, min_col=1, max_row=1, max_col=2)
    merged_cell_mock.parent.merged_cells.ranges = [merged_range_mock]
    top_left_cell_mock = mock_cell
    merged_cell_mock.parent.cell.return_value = top_left_cell_mock
    obj = ReaderHandler()
    obj.set_cell(merged_cell_mock, 100, "#,##0")
    assert top_left_cell_mock.value == 100
    assert top_left_cell_mock.number_format == "#,##0"

def test_set_cell_normal_case(workbook_and_worksheet):
    wb, ws = workbook_and_worksheet
    cell = ws.cell(row=1, column=1)
    helper = ReaderHandler()
    helper.set_cell(cell, 1000, "#,##0")
    assert cell.value == 1000
    assert cell.number_format == "#,##0"


def test_set_cell_with_merged_cells(workbook_and_worksheet):
    wb, ws = workbook_and_worksheet
    ws.merge_cells('A1:B2')
    cell = ws.cell(row=1, column=1) 
    helper = ReaderHandler()
    helper.set_cell(cell, 2000, "#,##0")
    top_left_cell = ws.cell(row=1, column=1)
    assert top_left_cell.value == 2000
    assert top_left_cell.number_format == "#,##0"